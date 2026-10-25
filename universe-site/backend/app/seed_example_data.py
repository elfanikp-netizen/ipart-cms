from pathlib import Path
import sys

from openpyxl import load_workbook
from sqlalchemy import select, text

from .auth import PortalUser
from .db import SessionLocal
from .models import Vehicle, Vin
from .routers.master_data import normalize_text, normalize_vin, normalized_key, resolve_vehicle, validate_year

EXAMPLE_WORKBOOK = Path(__file__).resolve().parents[2] / "data" / "Example Data.xlsx"


def import_example_data(path: Path = EXAMPLE_WORKBOOK) -> tuple[int, int, int]:
    if not path.is_file():
        raise FileNotFoundError(f"Place the workbook at {path} before running this import.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    db = SessionLocal()
    vehicle_count = vin_count = skipped = 0
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().casefold() for value in next(rows, ())]
        required = ("year", "make", "model", "vin")
        missing = [column for column in required if column not in headers]
        if missing:
            raise ValueError(f"Workbook is missing columns: {', '.join(missing)}")
        positions = {column: headers.index(column) for column in required}
        user_table = '"global"."users"' if db.bind and db.bind.dialect.name == "postgresql" else '"users"'
        portal_user = db.execute(text(f"SELECT id FROM {user_table} WHERE is_active = TRUE ORDER BY id LIMIT 1")).first()
        if portal_user is None:
            raise RuntimeError("No active Portal account is available to attribute imported records")
        importer = PortalUser(id=portal_user[0], is_active=True)
        with db.begin():
            for row_number, row in enumerate(rows, start=2):
                try:
                    year = validate_year(row[positions["year"]] if positions["year"] < len(row) else None)
                    make = normalize_text(row[positions["make"]] if positions["make"] < len(row) else None, "Make")
                    model = normalize_text(row[positions["model"]] if positions["model"] < len(row) else None, "Model")
                    vin = normalize_vin(row[positions["vin"]] if positions["vin"] < len(row) else None)
                    vehicle, created, _ = resolve_vehicle(db, year, make, model, importer.id)
                    vehicle_count += int(created)
                    existing = db.scalar(select(Vin).where(Vin.id_vehicle == vehicle.id, Vin.vin_key == vin))
                    if existing is None:
                        db.add(Vin(id_vehicle=vehicle.id, vin=vin, vin_key=vin, created_by=importer.id, updated_by=importer.id))
                        db.flush()
                        vin_count += 1
                    else:
                        skipped += 1
                except ValueError as error:
                    print(f"Skipping row {row_number}: {error}", file=sys.stderr)
                    skipped += 1
        return vehicle_count, vin_count, skipped
    finally:
        db.close()
        workbook.close()


if __name__ == "__main__":
    created_vehicles, created_vins, skipped_rows = import_example_data()
    print(f"Imported {created_vehicles} vehicle(s) and {created_vins} VIN(s); skipped {skipped_rows} unchanged/invalid row(s).")