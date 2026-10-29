from collections.abc import Sequence
from io import BytesIO
import math
import re

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import String, cast, func, or_, select, tuple_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..auth import PortalUser, get_current_user, get_db
from ..audit import log_activity
from ..models import Vehicle, Vin
from ..schemas import ImportOut, VehicleCreate, VehicleOut, VehicleUpdate, VinCreate, VinOut, VinUpdate

router = APIRouter(prefix="/api/universe", tags=["master data"])
DEFAULT_PAGE_SIZE = 10
MAX_PAGE_SIZE = 100
VIN_RE = re.compile(r"^[A-Z0-9]{5,32}$")


def normalize_text(value: object, label: str) -> str:
    text = " ".join(str(value or "").strip().split())
    if not text:
        raise ValueError(f"{label} is required")
    return text


def normalized_key(value: str) -> str:
    return " ".join(value.split()).casefold()


def normalize_vin(value: object) -> str:
    vin = normalize_text(value, "VIN").upper()
    if not VIN_RE.fullmatch(vin):
        raise ValueError("VIN must be 5-32 letters or digits")
    return vin


def find_duplicate_vin(db: Session, vehicle_id: int, vin: str, exclude_id: int | None = None):
    stmt = select(Vin.id).where(Vin.id_vehicle == vehicle_id, Vin.vin_key == vin)
    if exclude_id is not None:
        stmt = stmt.where(Vin.id != exclude_id)
    return db.scalar(stmt)


def validate_year(value: object) -> int:
    try:
        year_value = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError("Year must be a whole number") from error
    if not year_value.is_integer() or not 1886 <= year_value <= 2100:
        raise ValueError("Year must be a whole number between 1886 and 2100")
    return int(year_value)


def get_cell(row: Sequence, index: int) -> object:
    return row[index] if index < len(row) else None


def open_excel(file: UploadFile):
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    if suffix not in {"xlsx", "xlsm"}:
        raise HTTPException(400, "Upload an Excel workbook (.xlsx or .xlsm)")
    try:
        workbook = load_workbook(BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(400, "Could not read the uploaded workbook") from error
    rows = workbook.active.iter_rows(values_only=True)
    headers = [str(value or "").strip().casefold() for value in next(rows, ())]
    return workbook, rows, headers


def excel_indexes(headers: list[str], required: tuple[str, ...]) -> dict[str, int]:
    by_name = {header: index for index, header in enumerate(headers) if header}
    missing = [name for name in required if name.casefold() not in by_name]
    if missing:
        raise HTTPException(400, f"Missing required Excel columns: {', '.join(missing)}")
    return {name: by_name[name.casefold()] for name in required}


def resolve_vehicle(db: Session, year: int, make: str, model: str, user_id: int) -> tuple[Vehicle, bool, bool]:
    make_key, model_key = normalized_key(make), normalized_key(model)
    vehicle = db.scalar(select(Vehicle).where(
        Vehicle.year == year,
        Vehicle.make_key == make_key,
        Vehicle.model_key == model_key,
    ))
    if vehicle is None:
        vehicle = Vehicle(year=year, make=make, model=model, make_key=make_key, model_key=model_key,
                          created_by=user_id, updated_by=user_id)
        db.add(vehicle)
        db.flush()
        return vehicle, True, False
    changed = vehicle.make != make or vehicle.model != model
    if changed:
        vehicle.make = make
        vehicle.model = model
        vehicle.updated_by = user_id
    return vehicle, False, changed


def query_vehicles(db: Session, search: str, year: int | None, make: str):
    stmt = select(Vehicle)
    if year is not None:
        stmt = stmt.where(Vehicle.year == year)
    if make.strip():
        stmt = stmt.where(Vehicle.make_key == normalized_key(make))
    if search.strip():
        pattern = f"%{search.strip()}%"
        stmt = stmt.where(or_(
            Vehicle.make.ilike(pattern),
            Vehicle.model.ilike(pattern),
            cast(Vehicle.year, String).ilike(pattern),
        ))
    return stmt


def query_vins(db: Session, search: str, year: int | None, make: str, model: str):
    stmt = select(Vin).join(Vehicle)
    if year is not None:
        stmt = stmt.where(Vehicle.year == year)
    if make.strip():
        stmt = stmt.where(Vehicle.make_key == normalized_key(make))
    if model.strip():
        stmt = stmt.where(Vehicle.model_key == normalized_key(model))
    if search.strip():
        pattern = f"%{search.strip()}%"
        stmt = stmt.where(or_(
            Vin.vin.ilike(pattern),
            Vehicle.make.ilike(pattern),
            Vehicle.model.ilike(pattern),
            cast(Vehicle.year, String).ilike(pattern),
        ))
    return stmt


@router.get("/filters")
def list_filter_options(db: Session = Depends(get_db), _: PortalUser = Depends(get_current_user)):
    return {
        "years": list(db.scalars(select(Vehicle.year).distinct().order_by(Vehicle.year.desc())).all()),
        "makes": list(db.scalars(select(Vehicle.make).distinct().order_by(Vehicle.make)).all()),
        "models": list(db.scalars(select(Vehicle.model).distinct().order_by(Vehicle.model)).all()),
    }


@router.get("/vehicles/options")
def vehicle_options(
    search: str = "",
    limit: int = Query(50, ge=1, le=100),
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    stmt = select(Vehicle)
    if search.strip():
        term = f"%{search.strip()}%"
        stmt = stmt.where(or_(Vehicle.make.ilike(term), Vehicle.model.ilike(term),
                              cast(Vehicle.year, String).ilike(term)))
    rows = db.scalars(stmt.order_by(Vehicle.year.desc(), Vehicle.make_key, Vehicle.model_key).limit(limit)).all()
    return [{"id": row.id, "year": row.year, "make": row.make, "model": row.model,
             "label": f"{row.year} {row.make} {row.model}"} for row in rows]


def paginate(db: Session, stmt, page: int, page_size: int, order_by):
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
    items = db.scalars(stmt.order_by(*order_by).offset((page - 1) * page_size).limit(page_size)).all()
    return {"items": items, "page": page, "page_size": page_size, "total": total,
            "pages": math.ceil(total / page_size) if total else 0}


def excel_response(workbook: Workbook, filename: str):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/vehicles")
def list_vehicles(
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    search: str = "",
    year: int | None = None,
    make: str = "",
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    result = paginate(db, query_vehicles(db, search, year, make), page, page_size,
                      (Vehicle.year.desc(), Vehicle.make_key, Vehicle.model_key))
    ids = [vehicle.id for vehicle in result["items"]]
    counts = dict(db.execute(select(Vin.id_vehicle, func.count(Vin.id)).where(Vin.id_vehicle.in_(ids))
                             .group_by(Vin.id_vehicle)).all()) if ids else {}
    result["items"] = [VehicleOut.model_validate(vehicle).model_copy(update={"vin_count": counts.get(vehicle.id, 0)})
                       for vehicle in result["items"]]
    return result


@router.post("/vehicles", response_model=VehicleOut, status_code=201)
def create_vehicle(payload: VehicleCreate, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    make = normalize_text(payload.make, "Make")
    model = normalize_text(payload.model, "Model")
    vehicle, created, _ = resolve_vehicle(db, payload.year, make, model, user.id)
    if not created:
        raise HTTPException(409, "A vehicle with this Year, Make, and Model already exists")
    vehicle.notes = payload.notes.strip()
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "A vehicle with this Year, Make, and Model already exists") from error
    db.refresh(vehicle)
    log_activity(db, user, "create_vehicle", f"created {vehicle.year} {vehicle.make} {vehicle.model}")
    return VehicleOut.model_validate(vehicle)


@router.put("/vehicles/{vehicle_id}", response_model=VehicleOut)
def update_vehicle(vehicle_id: int, payload: VehicleUpdate, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    vehicle = db.get(Vehicle, vehicle_id)
    if vehicle is None:
        raise HTTPException(404, "Vehicle not found")
    make, model = normalize_text(payload.make, "Make"), normalize_text(payload.model, "Model")
    duplicate = db.scalar(select(Vehicle.id).where(
        Vehicle.year == payload.year,
        Vehicle.make_key == normalized_key(make),
        Vehicle.model_key == normalized_key(model),
        Vehicle.id != vehicle.id,
    ))
    if duplicate:
        raise HTTPException(409, "A vehicle with this Year, Make, and Model already exists")
    vehicle.year, vehicle.make, vehicle.model = payload.year, make, model
    vehicle.make_key, vehicle.model_key = normalized_key(make), normalized_key(model)
    vehicle.notes, vehicle.updated_by = payload.notes.strip(), user.id
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "A vehicle with this Year, Make, and Model already exists") from error
    db.refresh(vehicle)
    log_activity(db, user, "update_vehicle", f"updated {vehicle.year} {vehicle.make} {vehicle.model}")
    return VehicleOut.model_validate(vehicle)


@router.delete("/vehicles/{vehicle_id}", status_code=204)
def delete_vehicle(vehicle_id: int, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    vehicle = db.get(Vehicle, vehicle_id)
    if vehicle is None:
        raise HTTPException(404, "Vehicle not found")
    detail = f"deleted {vehicle.year} {vehicle.make} {vehicle.model}"
    db.delete(vehicle)
    db.commit()
    log_activity(db, user, "delete_vehicle", detail)


@router.get("/vehicles/export")
def export_vehicles(
    search: str = "", year: int | None = None, make: str = "",
    db: Session = Depends(get_db), _: PortalUser = Depends(get_current_user),
):
    rows = db.scalars(query_vehicles(db, search, year, make).order_by(Vehicle.year, Vehicle.make_key, Vehicle.model_key)).all()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Vehicle")
    sheet.append(["Year", "Make", "Model"])
    for vehicle in rows:
        sheet.append([vehicle.year, vehicle.make, vehicle.model])
    log_activity(db, _, "export_vehicles", f"exported {len(rows)} vehicle records")
    return excel_response(workbook, "vehicles.xlsx")


@router.post("/vehicles/import", response_model=ImportOut)
def import_vehicles(file: UploadFile = File(...), db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    workbook, rows, headers = open_excel(file)
    result = ImportOut()
    try:
        indexes = excel_indexes(headers, ("Year", "Make", "Model"))
        for row_number, row in enumerate(rows, start=2):
            try:
                with db.begin_nested():
                    year = validate_year(get_cell(row, indexes["Year"]))
                    make = normalize_text(get_cell(row, indexes["Make"]), "Make")
                    model = normalize_text(get_cell(row, indexes["Model"]), "Model")
                    _, created, updated = resolve_vehicle(db, year, make, model, user.id)
                    result.created += int(created)
                    result.updated += int(updated)
                    result.skipped += int(not created and not updated)
            except ValueError as error:
                result.skipped += 1
                if len(result.errors) < 50:
                    result.errors.append(f"Row {row_number}: {error}")
        db.commit()
        log_activity(db, user, "import_vehicles", f"created={result.created} updated={result.updated} skipped={result.skipped}")
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "Import conflicted with a duplicate vehicle") from error
    finally:
        workbook.close()
    return result


@router.get("/vins")
def list_vins(
    page: int = Query(1, ge=1), page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    search: str = "", year: int | None = None, make: str = "", model: str = "", vehicle_id: int | None = None,
    db: Session = Depends(get_db), _: PortalUser = Depends(get_current_user),
):
    stmt = query_vins(db, search, year, make, model)
    if vehicle_id is not None:
        stmt = stmt.where(Vin.id_vehicle == vehicle_id)
    result = paginate(db, stmt, page, page_size,
                      (Vehicle.year.desc(), Vehicle.make_key, Vehicle.model_key, Vin.vin_key))
    result["items"] = [VinOut.model_validate(item) for item in result["items"]]
    return result


@router.post("/vins", response_model=VinOut, status_code=201)
def create_vin(payload: VinCreate, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    vehicle = db.get(Vehicle, payload.id_vehicle)
    if vehicle is None:
        raise HTTPException(404, "Vehicle not found")
    vin = normalize_vin(payload.vin)
    if find_duplicate_vin(db, vehicle.id, vin):
        raise HTTPException(409, "This VIN already exists for the selected vehicle")
    item = Vin(id_vehicle=vehicle.id, vin=vin, vin_key=vin, notes=payload.notes.strip(),
               created_by=user.id, updated_by=user.id)
    db.add(item)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "This VIN already exists for the selected vehicle") from error
    db.refresh(item)
    log_activity(db, user, "create_vin", f"created VIN {item.vin}")
    return item


@router.put("/vins/{vin_id}", response_model=VinOut)
def update_vin(vin_id: int, payload: VinUpdate, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    item = db.get(Vin, vin_id)
    if item is None:
        raise HTTPException(404, "VIN not found")
    vehicle = db.get(Vehicle, payload.id_vehicle)
    if vehicle is None:
        raise HTTPException(404, "Vehicle not found")
    vin = normalize_vin(payload.vin)
    duplicate = find_duplicate_vin(db, vehicle.id, vin, exclude_id=item.id)
    if duplicate:
        raise HTTPException(409, "This VIN already exists for the selected vehicle")
    item.id_vehicle, item.vin, item.vin_key = vehicle.id, vin, vin
    item.notes, item.updated_by = payload.notes.strip(), user.id
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "This VIN already exists for the selected vehicle") from error
    db.refresh(item)
    log_activity(db, user, "update_vin", f"updated VIN {item.vin}")
    return item


@router.delete("/vins/{vin_id}", status_code=204)
def delete_vin(vin_id: int, db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    item = db.get(Vin, vin_id)
    if item is None:
        raise HTTPException(404, "VIN not found")
    detail = f"deleted VIN {item.vin}"
    db.delete(item)
    db.commit()
    log_activity(db, user, "delete_vin", detail)


@router.get("/vins/export")
def export_vins(
    search: str = "", year: int | None = None, make: str = "", model: str = "", vehicle_id: int | None = None,
    db: Session = Depends(get_db), _: PortalUser = Depends(get_current_user),
):
    stmt = query_vins(db, search, year, make, model)
    if vehicle_id is not None:
        stmt = stmt.where(Vin.id_vehicle == vehicle_id)
    rows = db.scalars(stmt
                      .order_by(Vehicle.year, Vehicle.make_key, Vehicle.model_key, Vin.vin_key)).all()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("VIN")
    sheet.append(["Year", "Make", "Model", "VIN"])
    for item in rows:
        sheet.append([item.vehicle.year, item.vehicle.make, item.vehicle.model, item.vin])
    log_activity(db, _, "export_vins", f"exported {len(rows)} VIN records")
    return excel_response(workbook, "vins.xlsx")


@router.post("/vins/export-by-vehicle-list")
def export_vins_by_vehicle_list(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    workbook, rows, headers = open_excel(file)
    try:
        indexes = excel_indexes(headers, ("Year", "Make", "Model"))
        keys: set[tuple[int, str, str]] = set()
        errors: list[str] = []
        for row_number, row in enumerate(rows, start=2):
            try:
                keys.add((
                    validate_year(get_cell(row, indexes["Year"])),
                    normalized_key(normalize_text(get_cell(row, indexes["Make"]), "Make")),
                    normalized_key(normalize_text(get_cell(row, indexes["Model"]), "Model")),
                ))
            except ValueError as error:
                if len(errors) < 20:
                    errors.append(f"Row {row_number}: {error}")
        if errors:
            raise HTTPException(400, "Invalid vehicle rows: " + "; ".join(errors))
    finally:
        workbook.close()

    matched_vehicles = db.scalars(select(Vehicle).where(
        tuple_(Vehicle.year, Vehicle.make_key, Vehicle.model_key).in_(keys)
    )).all() if keys else []
    vehicle_ids = [vehicle.id for vehicle in matched_vehicles]
    matched_vins = db.scalars(
        select(Vin).join(Vehicle).where(Vin.id_vehicle.in_(vehicle_ids))
        .order_by(Vehicle.year, Vehicle.make_key, Vehicle.model_key, Vin.vin_key)
    ).all() if vehicle_ids else []
    workbook_out = Workbook(write_only=True)
    sheet = workbook_out.create_sheet("VIN")
    sheet.append(["Year", "Make", "Model", "VIN"])
    for item in matched_vins:
        sheet.append([item.vehicle.year, item.vehicle.make, item.vehicle.model, item.vin])
    log_activity(db, user, "export_vins_by_vehicle_list", f"exported {len(matched_vins)} VIN records")
    return excel_response(workbook_out, "vins-by-vehicle-list.xlsx")


@router.post("/vins/import", response_model=ImportOut)
def import_vins(file: UploadFile = File(...), db: Session = Depends(get_db), user: PortalUser = Depends(get_current_user)):
    workbook, rows, headers = open_excel(file)
    result = ImportOut()
    seen_vins: set[tuple[int, str]] = set()
    try:
        indexes = excel_indexes(headers, ("Year", "Make", "Model", "VIN"))
        for row_number, row in enumerate(rows, start=2):
            try:
                with db.begin_nested():
                    year = validate_year(get_cell(row, indexes["Year"]))
                    make = normalize_text(get_cell(row, indexes["Make"]), "Make")
                    model = normalize_text(get_cell(row, indexes["Model"]), "Model")
                    vin = normalize_vin(get_cell(row, indexes["VIN"]))
                    vehicle, vehicle_created, vehicle_updated = resolve_vehicle(db, year, make, model, user.id)
                    result.vehicles_created += int(vehicle_created)
                    result.vehicles_updated += int(vehicle_updated)
                    duplicate_key = (vehicle.id, vin)
                    if duplicate_key in seen_vins:
                        result.skipped += 1
                        if len(result.errors) < 50:
                            result.errors.append(f"Row {row_number}: duplicate VIN {vin} in the uploaded workbook")
                        continue
                    seen_vins.add(duplicate_key)
                    existing = db.scalar(select(Vin).where(
                        Vin.id_vehicle == vehicle.id,
                        Vin.vin_key == vin,
                    ))
                    if existing:
                        changed = existing.vin != vin
                        existing.vin, existing.vin_key, existing.updated_by = vin, vin, user.id
                        result.updated += int(changed)
                        result.skipped += int(not changed)
                    else:
                        db.add(Vin(id_vehicle=vehicle.id, vin=vin, vin_key=vin, created_by=user.id, updated_by=user.id))
                        db.flush()
                        result.created += 1
            except ValueError as error:
                result.skipped += 1
                if len(result.errors) < 50:
                    result.errors.append(f"Row {row_number}: {error}")
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "Import conflicted with a duplicate VIN") from error
    finally:
        workbook.close()
    log_activity(
        db,
        user,
        "import_vins",
        f"vehicles_created={result.vehicles_created} vehicles_updated={result.vehicles_updated} "
        f"created={result.created} updated={result.updated} skipped={result.skipped}",
    )
    return result