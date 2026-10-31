import json
import math
from bisect import bisect_left
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import String, cast, func, or_, select, text, tuple_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..auth import PortalUser, get_current_user, get_db
from ..audit import log_activity
from ..models import CccMarketRecord
from ..schemas import CccMarketPageOut, CccMarketRecordOut
from .master_data import excel_indexes, get_cell, open_excel

router = APIRouter(prefix="/api/universe/market/ccc", tags=["CCC Market data"])


def parse_crawl_date(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        text = str(value or "").strip()
        if not text:
            raise ValueError("date_time is required")
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            parsed = None
            for date_format in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
                try:
                    parsed = datetime.strptime(text, date_format)
                    break
                except ValueError:
                    continue
            if parsed is None:
                raise ValueError("date_time must be a valid date or timestamp")
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def parse_market_data(value: Any) -> dict | list:
    if isinstance(value, (dict, list)):
        decoded = value
    else:
        text = str(value or "").strip()
        if not text:
            return []
        try:
            decoded = json.loads(text)
        except json.JSONDecodeError as error:
            raise ValueError("market_data must contain valid JSON") from error
    if not isinstance(decoded, (dict, list)):
        raise ValueError("market_data JSON must be an object or array")
    return decoded


def market_field(data: dict | list, field: str) -> str:
    values = data if isinstance(data, list) else [data]
    found = []
    for item in values:
        if not isinstance(item, dict):
            continue
        value = next((item[key] for key in item if key.casefold() == field), None)
        if value is not None and str(value).strip():
            rendered = str(value).strip()
            if rendered not in found:
                found.append(rendered)
    return ", ".join(found)


def market_entries(data: dict | list) -> list[dict]:
    entries = [data] if isinstance(data, dict) else data
    if any(not isinstance(entry, dict) for entry in entries):
        raise ValueError("Every market_data array entry must be a JSON object")
    return entries


def numeric_price(value: Any) -> Decimal | None:
    if value is None:
        return None
    rendered = str(value).strip().replace("$", "").replace(",", "")
    if not rendered:
        return None
    try:
        return Decimal(rendered)
    except InvalidOperation:
        return None


def filtered_ccc_records(db: Session, oem: str, vin: str, description: str, vendor: str, search: str):
    stmt = select(CccMarketRecord)
    if oem.strip():
        stmt = stmt.where(CccMarketRecord.oem == oem.strip())
    if vin.strip():
        stmt = stmt.where(CccMarketRecord.vin == vin.strip())
    if description.strip():
        stmt = stmt.where(CccMarketRecord.description.ilike(f"%{description.strip()}%"))
    if vendor.strip():
        stmt = stmt.where(CccMarketRecord.vendor.ilike(f"%{vendor.strip()}%"))
    if search.strip():
        pattern = f"%{search.strip()}%"
        stmt = stmt.where(or_(
            CccMarketRecord.oem.ilike(pattern),
            CccMarketRecord.vin.ilike(pattern),
            CccMarketRecord.description.ilike(pattern),
            CccMarketRecord.vendor.ilike(pattern),
            cast(CccMarketRecord.market_data, String).ilike(pattern),
        ))
    return stmt


@router.get("/filters")
def ccc_market_filters(db: Session = Depends(get_db), _: PortalUser = Depends(get_current_user)):
    rows = db.execute(select(
        CccMarketRecord.oem,
        CccMarketRecord.vin,
        CccMarketRecord.description,
        CccMarketRecord.vendor,
    ).distinct()).all()
    return {
        "oems": sorted({row.oem for row in rows if row.oem}),
        "vins": sorted({row.vin for row in rows if row.vin}),
        "descriptions": sorted({row.description for row in rows if row.description}),
        "vendors": sorted({row.vendor for row in rows if row.vendor}),
    }


@router.get("/offers")
def get_ccc_market_offers(
    oem: str = Query(..., min_length=1),
    vin: str = Query(..., min_length=1),
    crawl_date: datetime = Query(...),
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    rows = db.scalars(select(CccMarketRecord).where(
        CccMarketRecord.oem == oem,
        CccMarketRecord.vin == vin,
        CccMarketRecord.crawl_date == crawl_date,
    ).order_by(CccMarketRecord.market_index)).all()
    return [
        {"market_index": row.market_index, "market_data": row.market_data}
        for row in rows
        if row.market_data
    ]


@router.get("", response_model=CccMarketPageOut)
def list_ccc_market_records(
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    oem: str = "",
    vin: str = "",
    description: str = "",
    vendor: str = "",
    search: str = "",
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    stmt = filtered_ccc_records(db, oem, vin, description, vendor, search)
    groups = stmt.with_only_columns(
        CccMarketRecord.oem,
        CccMarketRecord.vin,
        CccMarketRecord.crawl_date,
    ).distinct().order_by(None).subquery()
    total = db.scalar(select(func.count()).select_from(groups)) or 0
    items = db.execute(select(groups).order_by(
        groups.c.crawl_date.desc(),
        groups.c.oem,
        groups.c.vin,
    ).offset((page - 1) * page_size).limit(page_size)).mappings().all()
    return {
        "items": [dict(item) for item in items],
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": math.ceil(total / page_size) if total else 0,
    }


@router.get("/export")
def export_ccc_market_records(
    oem: str = "",
    vin: str = "",
    description: str = "",
    vendor: str = "",
    search: str = "",
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    stmt = filtered_ccc_records(db, oem, vin, description, vendor, search)
    groups = stmt.with_only_columns(
        CccMarketRecord.oem,
        CccMarketRecord.vin,
        CccMarketRecord.crawl_date,
        func.count(CccMarketRecord.id).label("offer_count"),
    ).group_by(
        CccMarketRecord.oem,
        CccMarketRecord.vin,
        CccMarketRecord.crawl_date,
    ).order_by(
        CccMarketRecord.crawl_date.desc(),
        CccMarketRecord.oem,
        CccMarketRecord.vin,
    )
    rows = db.execute(groups).all()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("CCC Crawls")
    sheet.append(["OEM", "VIN", "Crawl Date", "Offer Count"])
    for row in rows:
        crawl_date = row.crawl_date.replace(tzinfo=None) if row.crawl_date.tzinfo else row.crawl_date
        sheet.append([row.oem, row.vin, crawl_date, row.offer_count])
    log_activity(db, user, "export_ccc_crawls", f"exported {len(rows)} filtered CCC crawls")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="ccc-crawls.xlsx"'},
    )


@router.get("/price-history")
def ccc_price_history(
    oem: str = Query(..., min_length=1),
    vin: str = Query(..., min_length=1),
    description: str = Query(..., min_length=1),
    from_date: date | None = None,
    to_date: date | None = None,
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    stmt = select(CccMarketRecord.crawl_date, CccMarketRecord.market_data).where(
        CccMarketRecord.oem == oem.strip(),
        CccMarketRecord.vin == vin.strip(),
    )
    if from_date is not None:
        stmt = stmt.where(CccMarketRecord.crawl_date >= datetime.combine(from_date, time.min, timezone.utc))
    if to_date is not None:
        stmt = stmt.where(CccMarketRecord.crawl_date < datetime.combine(to_date + timedelta(days=1), time.min, timezone.utc))

    prices: dict[tuple[str, datetime], tuple[str, Decimal]] = {}
    for crawl_date, market_data in db.execute(stmt):
        entries = [market_data] if isinstance(market_data, dict) else market_data
        if not isinstance(entries, list):
            continue
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            fields = {str(key).casefold(): value for key, value in entry.items()}
            offer_description = str(fields.get("description") or "").strip()
            vendor = str(fields.get("vendor") or "").strip()
            price = numeric_price(fields.get("ext-price"))
            if offer_description.casefold() != description.strip().casefold() or not vendor or price is None:
                continue
            key = (vendor.casefold(), crawl_date)
            previous = prices.get(key)
            if previous is None or price < previous[1]:
                prices[key] = (vendor, price)

    by_vendor: dict[str, list[tuple[datetime, Decimal]]] = defaultdict(list)
    vendor_labels: dict[str, str] = {}
    for (vendor_key, crawl_date), (vendor, price) in prices.items():
        vendor_labels[vendor_key] = vendor
        by_vendor[vendor_key].append((crawl_date, price))

    vendors = []
    for vendor_key, observations in by_vendor.items():
        observations.sort(key=lambda item: item[0])
        previous_price: Decimal | None = None
        history = []
        for crawl_date, price in observations:
            change = price - previous_price if previous_price is not None else None
            history.append({
                "crawl_date": crawl_date.isoformat(),
                "price": str(price),
                "previous_price": str(previous_price) if previous_price is not None else None,
                "change": str(change) if change is not None else None,
            })
            previous_price = price
        vendors.append({"vendor": vendor_labels[vendor_key], "observations": history})

    vendors.sort(key=lambda item: (
        Decimal(item["observations"][-1]["price"]),
        item["vendor"].casefold(),
    ))
    return {"oem": oem.strip(), "vin": vin.strip(), "description": description.strip(), "vendors": vendors}


def lowest_price_rows(
    db: Session,
    from_date: date | None,
    to_date: date | None,
    vehicle_keys: set[tuple[str, str]] | None = None,
    crawl_keys: set[tuple[str, str, datetime]] | None = None,
    target_vendor: str | None = None,
    target_description: str | None = None,
) -> list[dict[str, Any]]:
    if db.bind and db.bind.dialect.name == "postgresql":
        if (vehicle_keys is not None and not vehicle_keys) or (crawl_keys is not None and not crawl_keys):
            return []
        vehicle_cte = """
            requested_vehicle_keys AS (
                SELECT * FROM jsonb_to_recordset(CAST(:vehicle_keys AS jsonb)) AS key(oem text, vin text)
            ),
        """ if vehicle_keys is not None else ""
        crawl_cte = """
            requested_crawl_keys AS (
                SELECT * FROM jsonb_to_recordset(CAST(:crawl_keys AS jsonb)) AS key(oem text, vin text, crawl_date timestamptz)
            ),
        """ if crawl_keys is not None else ""
        target_group_cte = """
            target_crawl_groups AS (
                SELECT DISTINCT oem, vin, crawl_date
                FROM universe.ccc_market_data
                WHERE lower(btrim(vendor)) = lower(:target_vendor)
                  AND lower(btrim(description)) = lower(:target_description)
                  __TARGET_DATE_FILTERS__
            ),
        """ if target_vendor is not None and target_description is not None else ""
        vehicle_filter = """
            AND EXISTS (
                SELECT 1 FROM requested_vehicle_keys AS requested
                WHERE requested.oem = ccc_market_data.oem AND requested.vin = ccc_market_data.vin
            )
        """ if vehicle_keys is not None else ""
        crawl_filter = """
            AND EXISTS (
                SELECT 1 FROM requested_crawl_keys AS requested
                WHERE requested.oem = ccc_market_data.oem AND requested.vin = ccc_market_data.vin
                  AND requested.crawl_date = ccc_market_data.crawl_date
            )
        """ if crawl_keys is not None else ""
        target_group_filter = """
            AND EXISTS (
                SELECT 1 FROM target_crawl_groups AS target
                WHERE target.oem = ccc_market_data.oem AND target.vin = ccc_market_data.vin
                  AND target.crawl_date = ccc_market_data.crawl_date
            )
            AND lower(btrim(market_data ->> 'description')) = lower(:target_description)
        """ if target_vendor is not None and target_description is not None else ""
        description_filter = "AND lower(btrim(market_data ->> 'description')) = lower(:target_description)" if target_description is not None and target_vendor is None else ""
        date_filters = []
        if from_date is not None:
            date_filters.append("crawl_date >= :from_date")
        if to_date is not None:
            date_filters.append("crawl_date < :to_date")
        date_filter = " AND ".join(date_filters)
        if date_filter:
            date_filter = f"AND {date_filter}"
        target_date_filters = date_filter
        if target_vendor is not None and target_description is not None:
            target_group_cte = target_group_cte.replace("__TARGET_DATE_FILTERS__", target_date_filters)
        statement = text(f"""
            WITH {vehicle_cte}{crawl_cte}{target_group_cte}extracted AS (
                SELECT oem, vin, crawl_date,
                    lower(btrim(market_data ->> 'description')) AS description_key,
                    btrim(market_data ->> 'description') AS description,
                    btrim(market_data ->> 'vendor') AS vendor,
                    regexp_replace(COALESCE(market_data ->> 'ext-price', ''), '[$,[:space:]]', '', 'g') AS price_text
                FROM universe.ccc_market_data
                                WHERE TRUE {date_filter}
                        {vehicle_filter}
                                    {crawl_filter}
                                    {target_group_filter}
                          {description_filter}
                    ), vendor_prices AS (
                SELECT oem, vin, crawl_date, description_key,
                    min(description) AS description, vendor,
                    min(CASE WHEN price_text ~ '^[+-]?([0-9]+([.][0-9]*)?|[.][0-9]+)$'
                        THEN price_text::numeric END) AS price
                FROM extracted
                WHERE description_key <> '' AND vendor <> ''
                GROUP BY oem, vin, crawl_date, description_key, vendor
                HAVING min(CASE WHEN price_text ~ '^[+-]?([0-9]+([.][0-9]*)?|[.][0-9]+)$'
                    THEN price_text::numeric END) IS NOT NULL
            ), ranked AS (
                SELECT vendor_prices.*,
                    dense_rank() OVER (PARTITION BY oem, vin, crawl_date, description_key ORDER BY price) AS price_rank
                FROM vendor_prices
            ), price_groups AS (
                SELECT oem, vin, crawl_date, description_key,
                    min(price) FILTER (WHERE price_rank = 1) AS lowest_price,
                    min(price) FILTER (WHERE price_rank = 2) AS second_lowest_price,
                    string_agg(vendor, ', ' ORDER BY vendor) FILTER (WHERE price_rank = 1) AS lowest_vendors,
                    string_agg(vendor, ', ' ORDER BY vendor) FILTER (WHERE price_rank = 2) AS second_lowest_vendors
                FROM ranked
                GROUP BY oem, vin, crawl_date, description_key
            )
            SELECT r.oem, r.vin, r.crawl_date, r.vendor, r.description, r.price, g.lowest_price,
                g.lowest_vendors, g.second_lowest_price, g.second_lowest_vendors,
                CASE WHEN r.price_rank = 1 THEN 'won' ELSE 'lost' END AS outcome
            FROM ranked AS r
            JOIN price_groups AS g USING (oem, vin, crawl_date, description_key)
        """)
        params = {}
        if from_date is not None:
            params["from_date"] = datetime.combine(from_date, time.min, timezone.utc)
        if to_date is not None:
            params["to_date"] = datetime.combine(to_date + timedelta(days=1), time.min, timezone.utc)
        if vehicle_keys is not None:
            params["vehicle_keys"] = json.dumps([{"oem": oem, "vin": vin} for oem, vin in vehicle_keys])
        if crawl_keys is not None:
            params["crawl_keys"] = json.dumps([
                {"oem": oem, "vin": vin, "crawl_date": crawl_date.isoformat()}
                for oem, vin, crawl_date in crawl_keys
            ])
        if target_vendor is not None and target_description is not None:
            params["target_vendor"] = target_vendor.strip()
            params["target_description"] = target_description.strip()
        elif target_description is not None:
            params["target_description"] = target_description.strip()
        rows = db.execute(statement, params).mappings().all()
        return [dict(row) for row in rows]

    stmt = select(
        CccMarketRecord.oem,
        CccMarketRecord.vin,
        CccMarketRecord.crawl_date,
        CccMarketRecord.market_data,
    )
    if from_date is not None:
        stmt = stmt.where(CccMarketRecord.crawl_date >= datetime.combine(from_date, time.min, timezone.utc))
    if to_date is not None:
        stmt = stmt.where(CccMarketRecord.crawl_date < datetime.combine(to_date + timedelta(days=1), time.min, timezone.utc))
    if vehicle_keys is not None:
        if not vehicle_keys:
            return []
        stmt = stmt.where(tuple_(CccMarketRecord.oem, CccMarketRecord.vin).in_(vehicle_keys))

    crawls: dict[tuple[str, str, datetime], list[tuple[str, str, Decimal]]] = defaultdict(list)
    for oem, vin, crawl_date, offer in db.execute(stmt):
        if not isinstance(offer, dict):
            continue
        description = str(offer.get("description") or "").strip()
        vendor = str(offer.get("vendor") or "").strip()
        price = numeric_price(offer.get("ext-price"))
        if description and vendor and price is not None:
            crawls[(oem, vin, crawl_date)].append((description, vendor, price))

    result: list[dict[str, Any]] = []
    for crawl_key, offers in crawls.items():
        by_description: dict[str, list[tuple[str, str, Decimal]]] = defaultdict(list)
        for offer in offers:
            by_description[offer[0].casefold()].append(offer)
        for description_key, description_offers in by_description.items():
            lowest = min(offer[2] for offer in description_offers)
            description_label = description_offers[0][0]
            vendor_prices: dict[str, Decimal] = {}
            for _, vendor, price in description_offers:
                vendor_prices[vendor] = min(vendor_prices.get(vendor, price), price)
            second_lowest = min((price for price in vendor_prices.values() if price > lowest), default=None)
            lowest_vendors = sorted(vendor for vendor, price in vendor_prices.items() if price == lowest)
            second_lowest_vendors = sorted(
                vendor for vendor, price in vendor_prices.items() if second_lowest is not None and price == second_lowest
            )
            for vendor, price in vendor_prices.items():
                result.append({
                    "oem": crawl_key[0],
                    "vin": crawl_key[1],
                    "crawl_date": crawl_key[2],
                    "vendor": vendor,
                    "description": description_label,
                    "price": price,
                    "lowest_price": lowest,
                    "lowest_vendors": ", ".join(lowest_vendors),
                    "second_lowest_price": second_lowest,
                    "second_lowest_vendors": ", ".join(second_lowest_vendors),
                    "outcome": "won" if price == lowest else "lost",
                })
    return result


@router.get("/vendor-dashboard")
def ccc_vendor_dashboard(
    vendor: str = Query(..., min_length=1),
    from_date: date | None = None,
    to_date: date | None = None,
    description: str = "",
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    rows = [
        row for row in lowest_price_rows(
            db,
            from_date,
            to_date,
            target_vendor=vendor if description.strip() else None,
            target_description=description if description.strip() else None,
        )
        if row["vendor"].casefold() == vendor.strip().casefold()
        and (not description.strip() or row["description"].casefold() == description.strip().casefold())
    ]
    detail_rows = sorted(rows, key=lambda item: (item["crawl_date"], item["oem"], item["vin"]), reverse=True)[:250]
    detail_keys = {
        (row["oem"], row["vin"], row["description"].casefold(), row["crawl_date"])
        for row in detail_rows
    }
    history_rows: list[dict[str, Any]] = []
    if detail_keys and db.bind and db.bind.dialect.name == "postgresql":
        prior_records = db.execute(text("""
            WITH detail_keys AS (
                SELECT * FROM jsonb_to_recordset(CAST(:detail_keys AS jsonb))
                    AS key(oem text, vin text, description_key text, crawl_date timestamptz)
            ), prior_keys AS (
                SELECT DISTINCT keys.oem, keys.vin, keys.description_key, records.crawl_date
                FROM detail_keys AS keys
                JOIN universe.ccc_market_data AS records
                  ON records.oem = keys.oem AND records.vin = keys.vin
                 AND records.crawl_date < keys.crawl_date
                 AND lower(btrim(records.description)) = keys.description_key
            )
            SELECT records.oem, records.vin, records.crawl_date, records.market_data
            FROM universe.ccc_market_data AS records
            JOIN prior_keys AS prior
              ON prior.oem = records.oem AND prior.vin = records.vin
             AND prior.crawl_date = records.crawl_date
             AND prior.description_key = lower(btrim(records.description))
            WHERE btrim(records.vendor) <> ''
        """), {
            "detail_keys": json.dumps([
                {"oem": oem, "vin": vin, "description_key": description_key, "crawl_date": crawl_date.isoformat()}
                for oem, vin, description_key, crawl_date in detail_keys
            ]),
        }).mappings().all()
        grouped_prior: dict[tuple[str, str, datetime, str], list[tuple[str, Decimal]]] = defaultdict(list)
        for record in prior_records:
            offer = record["market_data"]
            if not isinstance(offer, dict):
                continue
            prior_vendor = str(offer.get("vendor") or "").strip()
            prior_price = numeric_price(offer.get("ext-price"))
            prior_description = str(offer.get("description") or "").strip().casefold()
            if prior_vendor and prior_price is not None:
                grouped_prior[(record["oem"], record["vin"], record["crawl_date"], prior_description)].append((prior_vendor, prior_price))
        for (prior_oem, prior_vin, prior_date, prior_description), offers in grouped_prior.items():
            prior_low = min(price for _, price in offers)
            prior_low_vendors = ", ".join(sorted(name for name, price in offers if price == prior_low))
            for prior_vendor, prior_price in offers:
                history_rows.append({
                    "oem": prior_oem,
                    "vin": prior_vin,
                    "description": prior_description,
                    "crawl_date": prior_date,
                    "vendor": prior_vendor,
                    "price": prior_price,
                    "lowest_vendors": prior_low_vendors,
                    "lowest_price": prior_low,
                })
    else:
        vehicle_keys = {(row["oem"], row["vin"]) for row in detail_rows}
        history_rows = lowest_price_rows(db, None, to_date, vehicle_keys)
    history_groups: dict[tuple[str, str, str, datetime], list[dict[str, Any]]] = defaultdict(list)
    for row in history_rows:
        history_groups[(row["oem"], row["vin"], row["description"].casefold(), row["crawl_date"])].append(row)

    previous_groups: dict[tuple[str, str, str], list[tuple[datetime, dict[str, Any]]]] = defaultdict(list)
    for key, group_rows in history_groups.items():
        lowest_row = group_rows[0]
        vendor_row = next((item for item in group_rows if item["vendor"].casefold() == vendor.strip().casefold()), None)
        previous_groups[key[:3]].append((key[3], {
            "lowest_vendors": lowest_row["lowest_vendors"],
            "lowest_price": lowest_row["lowest_price"],
            "vendor_price": vendor_row["price"] if vendor_row else None,
        }))
    for history in previous_groups.values():
        history.sort(key=lambda item: item[0])

    by_description: dict[str, dict[str, Any]] = {}
    by_day: dict[date, list[dict[str, Any]]] = defaultdict(list)
    total_gap = Decimal("0")
    wins = losses = 0
    for row in rows:
        won = row["outcome"] == "won"
        wins += int(won)
        losses += int(not won)
        total_gap += row["price"] - row["lowest_price"]
        description_key = row["description"].casefold()
        description_item = by_description.setdefault(description_key, {
            "description": row["description"], "wins": 0, "losses": 0,
        })
        description_item["wins"] += int(won)
        description_item["losses"] += int(not won)
        by_day[row["crawl_date"].date()].append(row)

    details = []
    for row in detail_rows:
        history = previous_groups.get((row["oem"], row["vin"], row["description"].casefold()), [])
        history_dates = [item[0] for item in history]
        previous_index = bisect_left(history_dates, row["crawl_date"]) - 1
        previous = history[previous_index] if previous_index >= 0 else None
        previous_data = previous[1] if previous else {}
        details.append({
            "oem": row["oem"],
            "vin": row["vin"],
            "description": row["description"],
            "crawl_date": row["crawl_date"].isoformat(),
            "price": str(row["price"]),
            "lowest_vendors": row["lowest_vendors"],
            "lowest_price": str(row["lowest_price"]),
            "gap": str(row["price"] - row["lowest_price"]),
            "outcome": row["outcome"],
            "previous_crawl_date": previous[0].isoformat() if previous else None,
            "previous_vendor_price": str(previous_data["vendor_price"]) if previous_data.get("vendor_price") is not None else None,
            "previous_lowest_vendors": previous_data.get("lowest_vendors", ""),
            "previous_lowest_price": str(previous_data["lowest_price"]) if previous_data.get("lowest_price") is not None else None,
        })

    descriptions = []
    for item in by_description.values():
        opportunities = item["wins"] + item["losses"]
        descriptions.append({
            **item,
            "opportunities": opportunities,
            "win_rate": round(item["wins"] / opportunities * 100, 1) if opportunities else 0,
        })
    descriptions.sort(key=lambda item: (-item["opportunities"], item["description"].casefold()))

    trend = []
    for crawl_day, day_rows in sorted(by_day.items()):
        count = Decimal(len(day_rows))
        average_price = sum((item["price"] for item in day_rows), Decimal("0")) / count
        average_lowest = sum((item["lowest_price"] for item in day_rows), Decimal("0")) / count
        trend.append({
            "date": crawl_day.isoformat(),
            "average_price": str(average_price.quantize(Decimal("0.01"))),
            "average_lowest_price": str(average_lowest.quantize(Decimal("0.01"))),
            "offers": len(day_rows),
        })

    return {
        "vendor": vendor.strip(),
        "from_date": from_date,
        "to_date": to_date,
        "priced_offers": len(rows),
        "wins": wins,
        "losses": losses,
        "win_rate": round(wins / len(rows) * 100, 1) if rows else 0,
        "average_gap": str((total_gap / Decimal(len(rows))).quantize(Decimal("0.01"))) if rows else "0.00",
        "trend": trend[-90:],
        "descriptions": descriptions,
        "details": details,
        "details_total": len(rows),
    }


@router.get("/lowest-price")
def ccc_lowest_price_summary(
    from_date: date | None = None,
    to_date: date | None = None,
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    rows = lowest_price_rows(db, from_date, to_date)
    vendor_descriptions: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    vendor_oems: dict[str, set[str]] = defaultdict(set)
    vendor_won_oems: dict[str, set[str]] = defaultdict(set)
    vendor_lost_oems: dict[str, set[str]] = defaultdict(set)
    vendor_wins: dict[str, set[tuple[str, str, datetime]]] = defaultdict(set)
    vendor_losses: dict[str, set[tuple[str, str, datetime]]] = defaultdict(set)
    description_oems: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    description_wins: dict[tuple[str, str, str], set[tuple[str, str, datetime]]] = defaultdict(set)
    description_losses: dict[tuple[str, str, str], set[tuple[str, str, datetime]]] = defaultdict(set)

    for row in rows:
        crawl_key = (row["oem"], row["vin"], row["crawl_date"])
        vendor = row["vendor"]
        description = row["description"]
        key = (vendor, description.casefold())
        vendor_descriptions[key].append(row)
        vendor_oems[vendor].add(row["oem"])
        description_key = (vendor, description.casefold(), row["outcome"])
        description_oems[description_key].add(row["oem"])
        if row["outcome"] == "won":
            vendor_wins[vendor].add(crawl_key)
            vendor_won_oems[vendor].add(row["oem"])
            description_wins[description_key].add(crawl_key)
        else:
            vendor_losses[vendor].add(crawl_key)
            vendor_lost_oems[vendor].add(row["oem"])
            description_losses[description_key].add(crawl_key)

    vendors = []
    for vendor in sorted(
        vendor_oems,
        key=lambda item: ("ipart" not in item.casefold(), -len(vendor_wins[item]), item.casefold()),
    ):
        description_keys = {key[1] for key in vendor_descriptions if key[0] == vendor}
        descriptions = []
        for description_key in description_keys:
            won_key = (vendor, description_key, "won")
            lost_key = (vendor, description_key, "lost")
            matching_rows = vendor_descriptions[(vendor, description_key)]
            description_label = matching_rows[0]["description"]
            descriptions.append({
                "description": description_label,
                "won_oem_count": len(description_oems[won_key]),
                "won_count": len(description_wins[won_key]),
                "lost_oem_count": len(description_oems[lost_key]),
                "lost_count": len(description_losses[lost_key]),
            })
        descriptions.sort(key=lambda item: (-(item["won_count"] + item["lost_count"]), item["description"].casefold()))
        vendors.append({
            "vendor": vendor,
            "oem_count": len(vendor_oems[vendor]),
            "won_oem_count": len(vendor_won_oems[vendor]),
            "lost_oem_count": len(vendor_lost_oems[vendor]),
            "lowest_count": len(vendor_wins[vendor]),
            "lost_count": len(vendor_losses[vendor]),
            "descriptions": descriptions,
        })

    return {
        "from_date": from_date,
        "to_date": to_date,
        "vendor_count": len(vendors),
        "oem_count": len({row["oem"] for row in rows}),
        "lowest_count": len({(row["oem"], row["vin"], row["crawl_date"]) for row in rows if row["outcome"] == "won"}),
        "lost_count": len({(row["oem"], row["vin"], row["crawl_date"]) for row in rows if row["outcome"] == "lost"}),
        "vendors": vendors,
    }


@router.get("/lowest-price/details")
def ccc_lowest_price_details(
    vendor: str = Query(..., min_length=1),
    description: str = Query(..., min_length=1),
    outcome: str = Query(..., pattern="^(won|lost)$"),
    metric: str = Query("crawl", pattern="^(oem|crawl)$"),
    from_date: date | None = None,
    to_date: date | None = None,
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    rows = [
        row for row in lowest_price_rows(
            db, from_date, to_date, target_vendor=vendor, target_description=description
        )
        if row["vendor"] == vendor
        and row["description"].casefold() == description.casefold()
        and row["outcome"] == outcome
    ]
    vehicle_keys = {(row["oem"], row["vin"]) for row in rows}
    all_rows = lowest_price_rows(db, None, None, vehicle_keys, target_description=description)
    previous_by_crawl: dict[tuple[str, str, str, datetime], dict[str, Any]] = {}
    for row in all_rows:
        key = (row["oem"], row["vin"], row["description"].casefold(), row["crawl_date"])
        existing = previous_by_crawl.get(key)
        if existing is None or row["lowest_price"] < existing["lowest_price"]:
            previous_by_crawl[key] = row

    all_crawls = sorted(previous_by_crawl.items(), key=lambda item: item[0][3])
    for row in rows:
        previous = [
            previous_row
            for (oem_key, vin_key, description_key, crawl_date), previous_row in all_crawls
            if oem_key == row["oem"]
            and vin_key == row["vin"]
            and description_key == row["description"].casefold()
            and crawl_date < row["crawl_date"]
        ]
        previous_row = max(previous, key=lambda item: item["crawl_date"], default=None)
        row["previous_crawl_date"] = previous_row["crawl_date"] if previous_row else None
        row["previous_lowest_vendors"] = previous_row["lowest_vendors"] if previous_row else ""
        row["previous_lowest_price"] = previous_row["lowest_price"] if previous_row else None
    rows.sort(key=lambda row: (row["crawl_date"], row["oem"], row["vin"]), reverse=True)
    if metric == "oem":
        by_oem: dict[str, dict[str, Any]] = {}
        for row in rows:
            item = by_oem.setdefault(row["oem"], {**row, "crawl_count": 0})
            item["crawl_count"] += 1
            if row["crawl_date"] > item["crawl_date"]:
                item.update(row)
                item["crawl_count"] = by_oem[row["oem"]]["crawl_count"]
        rows = sorted(by_oem.values(), key=lambda row: (row["oem"], row["vin"]))
    return [{
        **row,
        "price": str(row["price"]),
        "lowest_price": str(row["lowest_price"]),
        "second_lowest_price": str(row["second_lowest_price"]) if row["second_lowest_price"] is not None else None,
        "previous_lowest_price": str(row["previous_lowest_price"]) if row["previous_lowest_price"] is not None else None,
    } for row in rows]


@router.get("/lowest-price/export")
def export_ccc_lowest_price_details(
    vendor: str = Query(..., min_length=1),
    description: str = Query(..., min_length=1),
    outcome: str = Query(..., pattern="^(won|lost)$"),
    metric: str = Query("crawl", pattern="^(oem|crawl)$"),
    from_date: date | None = None,
    to_date: date | None = None,
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    rows = ccc_lowest_price_details(vendor, description, outcome, metric, from_date, to_date, db)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Lowest Price")
    sheet.append(["OEM", "VIN", "Crawl Date", "Vendor", "Description", "Outcome", "Price", "Lowest Vendor(s)", "Lowest Price", "Second Lowest Vendor(s)", "Second Lowest Price", "Previous Crawl Date", "Previous Lowest Vendor(s)", "Previous Lowest Price", "Crawls per OEM"])
    for row in rows:
        crawl_date = row["crawl_date"].replace(tzinfo=None) if row["crawl_date"].tzinfo else row["crawl_date"]
        previous_crawl_date = row["previous_crawl_date"]
        if previous_crawl_date is not None and previous_crawl_date.tzinfo:
            previous_crawl_date = previous_crawl_date.replace(tzinfo=None)
        sheet.append([row["oem"], row["vin"], crawl_date, row["vendor"], row["description"], row["outcome"], row["price"], row["lowest_vendors"], row["lowest_price"], row["second_lowest_vendors"], row["second_lowest_price"], previous_crawl_date, row["previous_lowest_vendors"], row["previous_lowest_price"], row.get("crawl_count", 1)])
    log_activity(db, user, "export_ccc_lowest_price_details", f"exported {len(rows)} CCC market detail rows")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    filename = f"ccc-{outcome}-{metric}-details.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/import")
def import_ccc_market_records(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    workbook, rows, headers = open_excel(file)
    created = updated = skipped = 0
    errors: list[str] = []
    try:
        indexes = excel_indexes(headers, ("VIN", "OEM", "Status", "date_time", "market_data"))
        for row_number, row in enumerate(rows, start=2):
            try:
                status = str(get_cell(row, indexes["Status"]) or "").strip()
                if status.casefold() != "ok":
                    skipped += 1
                    continue
                oem = str(get_cell(row, indexes["OEM"]) or "").strip().upper()
                vin = str(get_cell(row, indexes["VIN"]) or "").strip().upper()
                if not oem or not vin:
                    raise ValueError("OEM and VIN are required")
                crawl_date = parse_crawl_date(get_cell(row, indexes["date_time"]))
                market_data = parse_market_data(get_cell(row, indexes["market_data"]))
                entries = market_entries(market_data)
                stored_entries = entries or [{}]
                with db.begin_nested():
                    existing = db.scalars(select(CccMarketRecord).where(
                        CccMarketRecord.oem == oem,
                        CccMarketRecord.vin == vin,
                        CccMarketRecord.crawl_date == crawl_date,
                    ).order_by(CccMarketRecord.market_index)).all()
                    old_data = [record.market_data for record in existing]
                    if old_data == stored_entries:
                        skipped += len(stored_entries)
                    else:
                        if existing:
                            for record in existing:
                                db.delete(record)
                            db.flush()
                        for market_index, offer in enumerate(stored_entries):
                            db.add(CccMarketRecord(
                                oem=oem,
                                vin=vin,
                                crawl_date=crawl_date,
                                market_index=market_index,
                                description=market_field(offer, "description"),
                                vendor=market_field(offer, "vendor"),
                                market_data=offer,
                            ))
                        db.flush()
                        if existing:
                            updated += len(stored_entries)
                        else:
                            created += len(stored_entries)
            except (ValueError, TypeError) as error:
                skipped += 1
                if len(errors) < 50:
                    errors.append(f"Row {row_number}: {error}")
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "Import conflicted with an existing CCC market key") from error
    finally:
        workbook.close()
    log_activity(db, user, "import_ccc_market", f"created={created} updated={updated} skipped={skipped}")
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}
