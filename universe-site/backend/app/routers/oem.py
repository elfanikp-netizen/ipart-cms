import math
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import String, and_, cast, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from ..auth import PortalUser, get_current_user, get_db
from ..audit import log_activity
from ..models import Oem, Vehicle, Vin
from ..schemas import OemCreate, OemImportOut, OemOut, OemUpdate
from .master_data import (
    get_cell,
    normalize_text,
    normalize_vin,
    normalized_key,
    resolve_vehicle,
    validate_year,
)

router = APIRouter(prefix="/api/universe/oems", tags=["OEM master data"])
DEFAULT_PAGE_SIZE = 10
MAX_PAGE_SIZE = 100


def row_out(item: Oem) -> OemOut:
    return OemOut(
        id=item.id,
        id_vin=item.id_vin,
        oem=item.oem,
        description=item.description,
        notes=item.notes,
        vin=item.vin.vin,
        year=item.vin.vehicle.year,
        make=item.vin.vehicle.make,
        model=item.vin.vehicle.model,
        created_at=item.created_at,
        updated_at=item.updated_at,
    )


def query_oems(search: str):
    stmt = select(Oem).join(Oem.vin).join(Vin.vehicle)
    if search.strip():
        pattern = f"%{search.strip()}%"
        stmt = stmt.where(or_(
            Oem.oem.ilike(pattern),
            Oem.description.ilike(pattern),
            Vin.vin.ilike(pattern),
            Vehicle.make.ilike(pattern),
            Vehicle.model.ilike(pattern),
            cast(Vehicle.year, String).ilike(pattern),
        ))
    return stmt


def get_oem_or_404(db: Session, oem_id: int) -> Oem:
    item = db.scalar(
        select(Oem).options(joinedload(Oem.vin).joinedload(Vin.vehicle)).where(Oem.id == oem_id)
    )
    if item is None:
        raise HTTPException(404, "OEM record not found")
    return item


def duplicate_oem(db: Session, id_vin: int, oem_key: str, exclude_id: int | None = None) -> bool:
    stmt = select(Oem.id).where(Oem.id_vin == id_vin, Oem.oem_key == oem_key)
    if exclude_id is not None:
        stmt = stmt.where(Oem.id != exclude_id)
    return db.scalar(stmt) is not None


def normalize_oem(value: object) -> tuple[str, str]:
    oem = normalize_text(value, "OEM")
    if len(oem) > 150:
        raise ValueError("OEM must be 150 characters or fewer")
    return oem, normalized_key(oem)


@router.post("", response_model=OemOut, status_code=201)
def create_oem(
    payload: OemCreate,
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    vin_item = db.get(Vin, payload.id_vin)
    if vin_item is None:
        raise HTTPException(404, "VIN not found")
    try:
        oem, oem_key = normalize_oem(payload.oem)
    except ValueError as error:
        raise HTTPException(422, str(error)) from error
    if duplicate_oem(db, vin_item.id, oem_key):
        raise HTTPException(409, "This OEM already exists for the selected VIN")
    item = Oem(
        id_vin=vin_item.id,
        oem=oem,
        oem_key=oem_key,
        description=payload.description.strip(),
        notes=payload.notes.strip(),
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(item)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "This OEM already exists for the selected VIN") from error
    item = get_oem_or_404(db, item.id)
    log_activity(db, user, "create_oem", f"created OEM {item.oem} for VIN {item.vin.vin}")
    return row_out(item)


@router.put("/{oem_id}", response_model=OemOut)
def update_oem(
    oem_id: int,
    payload: OemUpdate,
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    item = get_oem_or_404(db, oem_id)
    vin_item = db.get(Vin, payload.id_vin)
    if vin_item is None:
        raise HTTPException(404, "VIN not found")
    try:
        oem, oem_key = normalize_oem(payload.oem)
    except ValueError as error:
        raise HTTPException(422, str(error)) from error
    if duplicate_oem(db, vin_item.id, oem_key, exclude_id=item.id):
        raise HTTPException(409, "This OEM already exists for the selected VIN")
    item.id_vin = vin_item.id
    item.oem = oem
    item.oem_key = oem_key
    item.description = payload.description.strip()
    item.notes = payload.notes.strip()
    item.updated_by = user.id
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "This OEM already exists for the selected VIN") from error
    item = get_oem_or_404(db, item.id)
    log_activity(db, user, "update_oem", f"updated OEM {item.oem} for VIN {item.vin.vin}")
    return row_out(item)


@router.delete("/{oem_id}", status_code=204)
def delete_oem(
    oem_id: int,
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    item = get_oem_or_404(db, oem_id)
    detail = f"deleted OEM {item.oem} for VIN {item.vin.vin}"
    db.delete(item)
    db.commit()
    log_activity(db, user, "delete_oem", detail)
    return Response(status_code=204)


def workbook_rows(file: UploadFile):
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    if suffix not in {"xlsx", "xlsm"}:
        raise HTTPException(400, "Upload an Excel workbook (.xlsx or .xlsm)")
    try:
        workbook = load_workbook(BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(400, "Could not read the uploaded workbook") from error
    rows = workbook.active.iter_rows(values_only=True)
    headers = [str(value or "").strip().casefold() for value in next(rows, ())]
    return workbook, rows, {header: index for index, header in enumerate(headers) if header}


def excel_response(workbook: Workbook, filename: str):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("")
def list_oems(
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    search: str = "",
    db: Session = Depends(get_db),
    _: PortalUser = Depends(get_current_user),
):
    stmt = query_oems(search)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
    items = db.scalars(
        stmt.options(joinedload(Oem.vin).joinedload(Vin.vehicle))
        .order_by(Oem.updated_at.desc(), Oem.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return {
        "items": [row_out(item) for item in items],
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": math.ceil(total / page_size) if total else 0,
    }


@router.get("/export")
def export_oems(
    search: str = "",
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    items = db.scalars(
        query_oems(search).options(joinedload(Oem.vin).joinedload(Vin.vehicle))
        .order_by(Oem.oem_key, Vin.vin_key)
    ).all()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("OEM")
    sheet.append(["OEM", "VIN", "Year", "Make", "Model", "Description", "Notes"])
    for item in items:
        vehicle = item.vin.vehicle
        sheet.append([item.oem, item.vin.vin, vehicle.year, vehicle.make, vehicle.model,
                      item.description, item.notes])
    log_activity(db, user, "export_oems", f"exported {len(items)} OEM records")
    return excel_response(workbook, "oems.xlsx")


@router.post("/export-by-list")
def export_oems_by_list(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    workbook, rows, indexes = workbook_rows(file)
    try:
        if "oem" not in indexes:
            raise HTTPException(400, "Missing required Excel column: OEM")
        keys: set[tuple[str, str] | tuple[str, None]] = set()
        for row_number, row in enumerate(rows, start=2):
            try:
                oem_key = normalized_key(normalize_text(get_cell(row, indexes["oem"]), "OEM"))
                vin_key = normalize_vin(get_cell(row, indexes["vin"])) if "vin" in indexes and get_cell(row, indexes["vin"]) else None
                keys.add((oem_key, vin_key))
            except ValueError as error:
                raise HTTPException(400, f"Row {row_number}: {error}") from error
    finally:
        workbook.close()

    all_oems = {oem_key for oem_key, vin_key in keys if vin_key is None}
    requested_vins = {(oem_key, vin_key) for oem_key, vin_key in keys if vin_key is not None}
    stmt = query_oems("").options(joinedload(Oem.vin).joinedload(Vin.vehicle))
    conditions = [Oem.oem_key.in_(all_oems)] if all_oems else []
    conditions.extend(
        and_(Oem.oem_key == oem_key, Vin.vin_key == vin_key)
        for oem_key, vin_key in requested_vins
    )
    stmt = stmt.where(or_(*conditions)) if conditions else stmt.where(False)
    items = db.scalars(stmt.order_by(Oem.oem_key, Vin.vin_key)).all()

    output = Workbook(write_only=True)
    sheet = output.create_sheet("OEM")
    sheet.append(["OEM", "VIN", "Year", "Make", "Model", "Description", "Notes"])
    for item in items:
        vehicle = item.vin.vehicle
        sheet.append([item.oem, item.vin.vin, vehicle.year, vehicle.make, vehicle.model,
                      item.description, item.notes])
    log_activity(db, user, "export_oems_by_list", f"exported {len(items)} OEM records from supplied list")
    return excel_response(output, "oems-by-list.xlsx")


@router.post("/import", response_model=OemImportOut)
def import_oems(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: PortalUser = Depends(get_current_user),
):
    workbook, rows, indexes = workbook_rows(file)
    result = OemImportOut()
    seen: set[tuple[int, str]] = set()
    try:
        required = ("oem", "vin")
        missing = [name.upper() for name in required if name not in indexes]
        if missing:
            raise HTTPException(400, f"Missing required Excel columns: {', '.join(missing)}")
        vehicle_columns = ("year", "make", "model")
        supplied = [name in indexes for name in vehicle_columns]
        if any(supplied) and not all(supplied):
            raise HTTPException(400, "Provide all three vehicle columns: Year, Make, and Model")

        for row_number, row in enumerate(rows, start=2):
            try:
                with db.begin_nested():
                    if "status" in indexes:
                        status = str(get_cell(row, indexes["status"]) or "").strip()
                        if status.casefold() != "ok":
                            result.skipped += 1
                            continue
                    oem = normalize_text(get_cell(row, indexes["oem"]), "OEM")
                    oem_key = normalized_key(oem)
                    vin_value = normalize_vin(get_cell(row, indexes["vin"]))
                    description_value = get_cell(row, indexes["description"]) if "description" in indexes else None
                    notes_value = get_cell(row, indexes["notes"]) if "notes" in indexes else None
                    description = " ".join(str(description_value or "").strip().split())
                    notes = " ".join(str(notes_value or "").strip().split())

                    if all(supplied):
                        year = validate_year(get_cell(row, indexes["year"]))
                        make = normalize_text(get_cell(row, indexes["make"]), "Make")
                        model = normalize_text(get_cell(row, indexes["model"]), "Model")
                        vehicle, vehicle_created, _ = resolve_vehicle(db, year, make, model, user.id)
                        result.vehicles_created += int(vehicle_created)
                        vin_item = db.scalar(select(Vin).where(
                            Vin.id_vehicle == vehicle.id,
                            Vin.vin_key == vin_value,
                        ))
                        if vin_item is None:
                            vin_item = Vin(id_vehicle=vehicle.id, vin=vin_value, vin_key=vin_value,
                                           created_by=user.id, updated_by=user.id)
                            db.add(vin_item)
                            db.flush()
                            result.vins_created += 1
                    else:
                        matches = db.scalars(select(Vin).where(Vin.vin_key == vin_value).limit(2)).all()
                        if len(matches) > 1:
                            raise ValueError(f"VIN {vin_value} matches multiple vehicles; include Year, Make, and Model")
                        if not matches:
                            raise ValueError(f"VIN {vin_value} was not found; include Year, Make, and Model to create it")
                        vin_item = matches[0]

                    duplicate_key = (vin_item.id, oem_key)
                    if duplicate_key in seen:
                        result.skipped += 1
                        if len(result.errors) < 50:
                            result.errors.append(f"Row {row_number}: duplicate OEM {oem} and VIN {vin_value} in workbook")
                        continue
                    seen.add(duplicate_key)
                    existing = db.scalar(select(Oem).where(
                        Oem.id_vin == vin_item.id,
                        Oem.oem_key == oem_key,
                    ))
                    if existing is None:
                        db.add(Oem(id_vin=vin_item.id, oem=oem, oem_key=oem_key,
                                   description=description, notes=notes,
                                   created_by=user.id, updated_by=user.id))
                        db.flush()
                        result.created += 1
                    else:
                        changed = (existing.oem, existing.description, existing.notes) != (oem, description, notes)
                        if changed:
                            existing.oem, existing.description, existing.notes = oem, description, notes
                            existing.updated_by = user.id
                        result.updated += int(changed)
                        result.skipped += int(not changed)
            except ValueError as error:
                result.skipped += 1
                if len(result.errors) < 50:
                    result.errors.append(f"Row {row_number}: {error}")
        db.commit()
        log_activity(db, user, "import_oems",
                     f"vehicles_created={result.vehicles_created} vins_created={result.vins_created} "
                     f"created={result.created} updated={result.updated} skipped={result.skipped}")
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "Import conflicted with a duplicate OEM record") from error
    finally:
        workbook.close()
    return result
